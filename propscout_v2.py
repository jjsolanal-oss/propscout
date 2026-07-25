#!/usr/bin/env python3
"""
PropScout V2 Phase 1 — Scraper-first architecture
Scrapes Brevitas, LandBrokerMLS, CommercialExchange, Catylist, Land.com
for San Antonio TX metro flex-industrial, small-bay-industrial, and land.

One Claude API call classifies + flags all deduplicated listings.
"""

import anthropic
import argparse
import json
import os
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from typing import Optional
from urllib.parse import urljoin, urlparse
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────────────────
#  COST GUARD  ── abort before any second Claude call
# ─────────────────────────────────────────────────────────────────────────────

MAX_API_CALLS   = 1
_api_calls_used = 0


def _use_api_call(label: str = "classify") -> None:
    global _api_calls_used
    if _api_calls_used >= MAX_API_CALLS:
        sys.exit(
            f"\n  ABORT: MAX_API_CALLS={MAX_API_CALLS} already reached. "
            f"Refused to make another call ({label}). "
            "Check code — only one Claude call is allowed."
        )
    _api_calls_used += 1


# ─────────────────────────────────────────────────────────────────────────────
#  CONFIGURATION
# ─────────────────────────────────────────────────────────────────────────────

MODEL         = "claude-sonnet-4-6"
RAW_JSON      = "listings_raw.json"
EXCEL_PATH    = "propscout_sa_phase1.xlsx"
REQUEST_DELAY = 2.0     # seconds between HTTP requests
MAX_PAGES     = 5       # pagination cap per source
REQ_TIMEOUT   = 20      # HTTP timeout in seconds
CURRENT_YEAR  = datetime.now().year

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0.0.0 Safari/537.36"
)

METRO_CITIES = frozenset({
    "san antonio", "schertz", "new braunfels", "converse", "universal city",
    "live oak", "helotes", "boerne", "seguin", "san marcos",
})

VA_PSF_THRESHOLD  = 120.0
DEV_PSF_THRESHOLD = 5.0
SQF_PER_ACRE      = 43_560

# Toggle sources without editing scraper code
SOURCES: dict[str, bool] = {
    "Brevitas":           True,
    "LandBrokerMLS":      True,
    "CommercialExchange": True,
    "Catylist":           True,
    "Land.com":           True,
}


# ─────────────────────────────────────────────────────────────────────────────
#  DATA MODELS
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class Listing:
    listing_name:     str            = ""
    address:          str            = ""
    city:             str            = ""
    state:            str            = "TX"
    zip_code:         str            = ""
    price:            Optional[float] = None
    total_sqft:       Optional[float] = None
    price_per_sqft:   Optional[float] = None
    year_built:       Optional[int]   = None
    lot_size_acres:   Optional[float] = None
    zoning:           Optional[str]   = None
    days_on_market:   Optional[int]   = None
    property_type:    str            = ""
    listing_source:   str            = ""
    also_listed_on:   str            = ""
    listing_url:      str            = ""
    deal_type:        str            = ""
    deal_type_reason: str            = ""
    flag:             str            = ""


@dataclass
class ScraperResult:
    source:         str
    http_status:    Optional[int]  = None
    listings_found: int            = 0
    error:          Optional[str]  = None
    listings:       list[Listing]  = field(default_factory=list)


# ─────────────────────────────────────────────────────────────────────────────
#  NUMBER / TEXT PARSING HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _parse_price(text) -> Optional[float]:
    if text is None:
        return None
    s = re.sub(r"[,$\s]", "", str(text))
    m = re.search(r"([\d]+\.?\d*)\s*([KkMmBb])?", s)
    if not m:
        return None
    v = float(m.group(1))
    suffix = (m.group(2) or "").upper()
    v *= {"K": 1e3, "M": 1e6, "B": 1e9}.get(suffix, 1)
    return v if v > 0 else None


def _parse_sqft(text) -> Optional[float]:
    if text is None:
        return None
    s = re.sub(r"[,\s]*(sq\.?\s*ft\.?|sqft|sf)\b", "", str(text), flags=re.IGNORECASE)
    m = re.search(r"[\d]+\.?\d*", s.replace(",", ""))
    if m:
        try:
            return float(m.group())
        except ValueError:
            pass
    return None


def _parse_int(text) -> Optional[int]:
    if text is None:
        return None
    m = re.search(r"\d+", str(text).replace(",", ""))
    return int(m.group()) if m else None


def _parse_acres(text) -> Optional[float]:
    if text is None:
        return None
    s = re.sub(r"[,\s]*(acres?|ac\.?)\b", "", str(text), flags=re.IGNORECASE)
    m = re.search(r"[\d]+\.?\d*", s.replace(",", ""))
    if m:
        v = float(m.group())
        return v if v > 0 else None
    return None


def _infer_prop_type(text: str) -> str:
    t = text.lower()
    if any(w in t for w in ("land", " lot ", "acreage", "tract", "vacant land", "raw land")):
        return "land"
    if any(w in t for w in ("small bay", "multi-tenant", "multitenant")):
        return "small bay industrial"
    if any(w in t for w in ("flex", "industrial", "warehouse", "distribution", "manufacturing")):
        return "flex industrial"
    return ""


def _in_metro(city: str) -> bool:
    return not city or city.strip().lower() in METRO_CITIES


def _parse_address(raw: str) -> tuple[str, str, str, str]:
    """Return (street, city, state, zip)."""
    raw = str(raw).strip()
    zm = re.search(r"\b(\d{5})\b", raw)
    zip_code = zm.group(1) if zm else ""

    parts = [p.strip() for p in re.split(r",", raw)]
    if len(parts) >= 3:
        street = parts[0]
        city   = parts[1]
    elif len(parts) == 2:
        street = parts[0]
        cm = re.match(r"(.+?)\s+(TX|Texas)\b", parts[1], re.IGNORECASE)
        city = cm.group(1).strip() if cm else parts[1]
    else:
        street = raw
        city   = ""

    city = re.sub(r"\s+(TX|Texas)\b", "", city, flags=re.IGNORECASE).strip()
    city = re.sub(r"\s*\d{5}\b", "", city).strip()
    return street, city, "TX", zip_code


# ─────────────────────────────────────────────────────────────────────────────
#  ROBOTS.TXT  — respect site crawling policies
# ─────────────────────────────────────────────────────────────────────────────

_robots_cache: dict[str, Optional[RobotFileParser]] = {}


def _robots_allows(url: str) -> bool:
    parsed  = urlparse(url)
    base    = f"{parsed.scheme}://{parsed.netloc}"
    if base not in _robots_cache:
        rp = RobotFileParser()
        rp.set_url(f"{base}/robots.txt")
        try:
            rp.read()
            _robots_cache[base] = rp
        except Exception:
            _robots_cache[base] = None   # unreadable → assume allowed
    rp = _robots_cache[base]
    return rp is None or rp.can_fetch(USER_AGENT, url)


# ─────────────────────────────────────────────────────────────────────────────
#  SHARED HTML EXTRACTION HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _json_ld_listings(soup: BeautifulSoup, source: str) -> list[Listing]:
    out = []
    for tag in soup.find_all("script", type="application/ld+json"):
        try:
            data = json.loads(tag.string or "")
        except Exception:
            continue
        items: list = []
        if isinstance(data, dict):
            if data.get("@type") == "ItemList":
                items = data.get("itemListElement", [])
            elif data.get("@type") in ("RealEstateListing", "Product"):
                items = [data]
        elif isinstance(data, list):
            items = data
        for item in items:
            if not isinstance(item, dict):
                continue
            thing  = item.get("item", item)
            name   = str(thing.get("name", "")).strip()
            ao     = thing.get("address", {})
            street = str(ao.get("streetAddress", "") if isinstance(ao, dict) else ao).strip()
            city   = str(ao.get("addressLocality", "") if isinstance(ao, dict) else "").strip()
            zp     = str(ao.get("postalCode", "") if isinstance(ao, dict) else "").strip()
            price  = _parse_price(thing.get("price") or
                                  (thing.get("offers") or {}).get("price"))
            url    = str(thing.get("url", "")).strip()
            if not name and not street:
                continue
            out.append(Listing(
                listing_name  = name[:120],
                address       = street[:120],
                city          = city,
                state         = "TX",
                zip_code      = zp[:10],
                price         = price,
                property_type = _infer_prop_type(name + " " + street),
                listing_source = source,
                listing_url   = url[:300],
            ))
    return out


def _next_data_listings(soup: BeautifulSoup, source: str) -> list[Listing]:
    tag = soup.find("script", id="__NEXT_DATA__")
    if not tag or not tag.string:
        return []
    try:
        data = json.loads(tag.string)
    except Exception:
        return []

    pp = data.get("props", {}).get("pageProps", {})
    raw: Optional[list] = None
    for key in ("listings", "properties", "results", "data", "items",
                "searchResults", "propertyList"):
        if isinstance(pp.get(key), list):
            raw = pp[key]
            break
    if raw is None:
        raw = _deep_find_list(data)
    if not raw:
        return []

    out = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        addr  = str(item.get("address") or item.get("streetAddress") or
                    item.get("street") or "").strip()
        city  = str(item.get("city") or item.get("cityName") or "").strip()
        zp    = str(item.get("zip") or item.get("zipCode") or "").strip()
        name  = str(item.get("name") or item.get("title") or
                    item.get("listingTitle") or addr or "Listing").strip()
        price = _parse_price(item.get("price") or item.get("askingPrice") or
                             item.get("listPrice"))
        sqft  = _parse_sqft(item.get("sqft") or item.get("buildingSize") or
                            item.get("totalSize"))
        year  = _parse_int(item.get("yearBuilt") or item.get("year_built"))
        acres = _parse_acres(item.get("lotSize") or item.get("acreage"))
        dom   = _parse_int(item.get("daysOnMarket") or item.get("days_on_market"))
        slug  = item.get("slug") or item.get("id") or item.get("propertyId") or ""
        url   = str(item.get("url") or item.get("listingUrl") or
                    (f"/{slug}" if slug else "")).strip()
        psf   = round(price / sqft, 2) if (price and sqft and sqft > 0) else None
        if not addr and not name:
            continue
        out.append(Listing(
            listing_name   = name[:120],
            address        = addr[:120],
            city           = city,
            state          = "TX",
            zip_code       = zp[:10],
            price          = price,
            total_sqft     = sqft,
            price_per_sqft = psf,
            year_built     = year,
            lot_size_acres = acres,
            days_on_market = dom,
            property_type  = _infer_prop_type(name),
            listing_source = source,
            listing_url    = url[:300],
        ))
    return out


def _deep_find_list(obj, depth: int = 0) -> Optional[list]:
    if depth > 7:
        return None
    if isinstance(obj, list) and obj and isinstance(obj[0], dict):
        keys = set(obj[0].keys())
        if keys & {"address", "price", "askingPrice", "listPrice",
                   "streetAddress", "sqft", "buildingSize", "slug"}:
            return obj
    if isinstance(obj, dict):
        for v in obj.values():
            r = _deep_find_list(v, depth + 1)
            if r:
                return r
    return None


def _safe_get(session: requests.Session,
              url: str, result: ScraperResult) -> Optional[requests.Response]:
    """GET with robots check, error capture, and status recording."""
    if not _robots_allows(url):
        result.error = "robots.txt disallows crawling"
        return None
    try:
        resp = session.get(url, timeout=REQ_TIMEOUT)
    except requests.RequestException as exc:
        result.error = f"Request error: {exc}"
        return None
    result.http_status = resp.status_code
    if resp.status_code == 403:
        result.error = "HTTP 403 – blocked"
        return None
    if resp.status_code == 429:
        result.error = "HTTP 429 – rate limited"
        return None
    if resp.status_code not in (200, 301, 302):
        result.error = f"HTTP {resp.status_code}"
        return None
    if ("cf-ray" in resp.headers or
            "checking your browser" in resp.text.lower()):
        result.error = "Cloudflare JS challenge"
        result.http_status = 403
        return None
    return resp


# ─────────────────────────────────────────────────────────────────────────────
#  SCRAPER 1 — Brevitas
# ─────────────────────────────────────────────────────────────────────────────

def scrape_brevitas(session: requests.Session) -> ScraperResult:
    """Brevitas.com — commercial real estate marketplace."""
    result = ScraperResult(source="Brevitas")
    listings: list[Listing] = []

    search_urls = [
        "https://brevitas.com/buy?q=San+Antonio+TX+industrial",
        "https://brevitas.com/buy?q=San+Antonio+TX+land",
        "https://brevitas.com/search?q=San+Antonio+TX&property_type=Industrial",
    ]

    for url in search_urls:
        resp = _safe_get(session, url, result)
        if resp is None:
            if result.error:
                return result
            continue

        soup = BeautifulSoup(resp.text, "html.parser")

        # Strategy 1: JSON-LD
        listings.extend(_json_ld_listings(soup, "Brevitas"))

        # Strategy 2: Next.js __NEXT_DATA__
        listings.extend(_next_data_listings(soup, "Brevitas"))

        # Strategy 3: listing cards
        for card in soup.select(
            ".listing-card, .property-card, [data-testid='listing'], "
            ".listing, article.property, .search-result"
        ):
            try:
                name_el  = card.select_one("h2, h3, h4, .title, .name, .listing-title")
                addr_el  = card.select_one(".address, .location, [class*='address']")
                price_el = card.select_one(".price, [class*='price'], .asking-price")
                size_el  = card.select_one(".sqft, .size, [class*='sqft'], [class*='size']")
                link     = card.select_one("a[href]")

                name  = name_el.get_text(strip=True)  if name_el  else ""
                addr  = addr_el.get_text(strip=True)  if addr_el  else ""
                price = _parse_price(price_el.get_text() if price_el else None)
                sqft  = _parse_sqft(size_el.get_text()  if size_el  else None)
                href  = link["href"] if link else ""
                if href and not href.startswith("http"):
                    href = f"https://brevitas.com{href}"

                street, city, _, zp = _parse_address(addr)
                psf = round(price / sqft, 2) if (price and sqft and sqft > 0) else None

                if name or street:
                    listings.append(Listing(
                        listing_name   = name,
                        address        = street,
                        city           = city,
                        state          = "TX",
                        zip_code       = zp,
                        price          = price,
                        total_sqft     = sqft,
                        price_per_sqft = psf,
                        property_type  = _infer_prop_type(name + " " + addr),
                        listing_source = "Brevitas",
                        listing_url    = href,
                    ))
            except Exception:
                continue

        time.sleep(REQUEST_DELAY)

    listings = [l for l in listings if _in_metro(l.city)]
    result.listings       = listings
    result.listings_found = len(listings)
    if not listings and not result.error:
        result.error = "No listings found (selectors may need updating)"
    return result


# ─────────────────────────────────────────────────────────────────────────────
#  SCRAPER 2 — LandBrokerMLS
# ─────────────────────────────────────────────────────────────────────────────

def scrape_landbrokermls(session: requests.Session) -> ScraperResult:
    """LandBrokerMLS.com — broker MLS for rural / ranch / land listings."""
    result = ScraperResult(source="LandBrokerMLS")
    listings: list[Listing] = []

    base_urls = [
        "https://www.landbrokermls.com/listings/?state_name=Texas&q=San+Antonio",
        "https://www.landbrokermls.com/listings/?state_name=Texas&q=Bexar+County",
        "https://www.landbrokermls.com/listings/?state_name=Texas&county=Bexar",
    ]

    for base_url in base_urls:
        for page in range(1, MAX_PAGES + 1):
            url = base_url if page == 1 else f"{base_url}&page={page}"
            resp = _safe_get(session, url, result)
            if resp is None:
                break

            soup = BeautifulSoup(resp.text, "html.parser")

            # JSON-LD / embedded JSON first
            listings.extend(_json_ld_listings(soup, "LandBrokerMLS"))
            listings.extend(_next_data_listings(soup, "LandBrokerMLS"))

            # HTML cards — LandBrokerMLS uses a standard WordPress-style layout
            cards = soup.select(
                ".listing-item, .property-listing, .listing-card, "
                ".idx-listing, article.listing, .search-item"
            )
            page_hits = 0
            for card in cards:
                try:
                    name_el  = card.select_one("h2, h3, .listing-title, .property-name")
                    addr_el  = card.select_one(".listing-address, .address, .location")
                    price_el = card.select_one(".listing-price, .price, .asking-price")
                    size_el  = card.select_one(".listing-acres, .acreage, .size, .acres")
                    link     = card.select_one("a[href]")

                    name  = name_el.get_text(strip=True)  if name_el  else ""
                    addr  = addr_el.get_text(strip=True)  if addr_el  else ""
                    price = _parse_price(price_el.get_text() if price_el else None)
                    acres = _parse_acres(size_el.get_text()  if size_el  else None)
                    href  = link["href"] if link else ""
                    if href and not href.startswith("http"):
                        href = urljoin("https://www.landbrokermls.com", href)

                    street, city, _, zp = _parse_address(addr)
                    psf = None
                    if price and acres and acres > 0:
                        psf = round(price / (acres * SQF_PER_ACRE), 2)

                    if name or street:
                        listings.append(Listing(
                            listing_name   = name,
                            address        = street,
                            city           = city,
                            state          = "TX",
                            zip_code       = zp,
                            price          = price,
                            lot_size_acres = acres,
                            price_per_sqft = psf,
                            property_type  = "land",
                            listing_source = "LandBrokerMLS",
                            listing_url    = href,
                        ))
                        page_hits += 1
                except Exception:
                    continue

            time.sleep(REQUEST_DELAY)
            if page_hits == 0:
                break   # no more pages

        if listings:
            break   # got results from this base_url, no need to try others

    listings = [l for l in listings if _in_metro(l.city)]
    result.listings       = listings
    result.listings_found = len(listings)
    if not listings and not result.error:
        result.error = "No listings found"
    return result


# ─────────────────────────────────────────────────────────────────────────────
#  SCRAPER 3 — CommercialExchange
# ─────────────────────────────────────────────────────────────────────────────

def scrape_commercialexchange(session: requests.Session) -> ScraperResult:
    """CommercialExchange.com — Catylist-powered commercial search."""
    result   = ScraperResult(source="CommercialExchange")
    listings: list[Listing] = []

    search_urls = [
        ("flex industrial", "https://www.commercialexchange.com/commercial-real-estate/san-antonio-tx/industrial/for-sale/"),
        ("land",            "https://www.commercialexchange.com/commercial-real-estate/san-antonio-tx/land/for-sale/"),
        ("flex industrial", "https://www.commercialexchange.com/search?location=San+Antonio%2C+TX&propertyType=industrial"),
        ("land",            "https://www.commercialexchange.com/search?location=San+Antonio%2C+TX&propertyType=land"),
    ]

    tried = set()
    for prop_type, url in search_urls:
        if url in tried:
            continue
        tried.add(url)

        resp = _safe_get(session, url, result)
        if resp is None:
            if result.error:
                return result
            continue

        soup = BeautifulSoup(resp.text, "html.parser")

        listings.extend(_json_ld_listings(soup, "CommercialExchange"))
        listings.extend(_next_data_listings(soup, "CommercialExchange"))

        for card in soup.select(
            ".listing-card, .property-card, .search-result, "
            ".property-row, article.property, [data-testid*='listing']"
        ):
            try:
                name_el  = card.select_one("h2, h3, h4, .name, .title")
                addr_el  = card.select_one(".address, .property-address, .location")
                price_el = card.select_one(".price, .asking-price, [class*='price']")
                size_el  = card.select_one(".sqft, .size, [class*='sqft'], [class*='size']")
                link     = card.select_one("a[href]")

                name  = name_el.get_text(strip=True)  if name_el  else ""
                addr  = addr_el.get_text(strip=True)  if addr_el  else ""
                price = _parse_price(price_el.get_text() if price_el else None)
                sqft  = _parse_sqft(size_el.get_text()  if size_el  else None)
                href  = link["href"] if link else ""
                if href and not href.startswith("http"):
                    href = urljoin("https://www.commercialexchange.com", href)

                street, city, _, zp = _parse_address(addr)
                psf = round(price / sqft, 2) if (price and sqft and sqft > 0) else None

                if name or street:
                    listings.append(Listing(
                        listing_name   = name,
                        address        = street,
                        city           = city,
                        state          = "TX",
                        zip_code       = zp,
                        price          = price,
                        total_sqft     = sqft,
                        price_per_sqft = psf,
                        property_type  = prop_type,
                        listing_source = "CommercialExchange",
                        listing_url    = href,
                    ))
            except Exception:
                continue

        time.sleep(REQUEST_DELAY)

    listings = [l for l in listings if _in_metro(l.city)]
    result.listings       = listings
    result.listings_found = len(listings)
    if not listings and not result.error:
        result.error = "No listings found"
    return result


# ─────────────────────────────────────────────────────────────────────────────
#  SCRAPER 4 — Catylist
# ─────────────────────────────────────────────────────────────────────────────

def scrape_catylist(session: requests.Session) -> ScraperResult:
    """Catylist.com — commercial real estate search."""
    result   = ScraperResult(source="Catylist")
    listings: list[Listing] = []

    search_urls = [
        "https://www.catylist.com/san-antonio-tx-commercial-real-estate/industrial/for-sale/",
        "https://www.catylist.com/san-antonio-tx-commercial-real-estate/land/for-sale/",
        "https://www.catylist.com/search?location=san-antonio-tx&type=sale&propertyType=industrial",
        "https://www.catylist.com/search?location=san-antonio-tx&type=sale&propertyType=land",
    ]

    tried: set[str] = set()
    for url in search_urls:
        if url in tried:
            continue
        tried.add(url)

        resp = _safe_get(session, url, result)
        if resp is None:
            if result.error:
                return result
            continue

        soup = BeautifulSoup(resp.text, "html.parser")

        listings.extend(_json_ld_listings(soup, "Catylist"))
        listings.extend(_next_data_listings(soup, "Catylist"))

        for card in soup.select(
            ".listing-card, .property-card, .search-listing, "
            ".listing-item, article, .result-item"
        ):
            try:
                name_el  = card.select_one("h2, h3, h4, .title, .listing-title")
                addr_el  = card.select_one(".address, .location, .listing-address")
                price_el = card.select_one(".price, .asking-price, [class*='price']")
                size_el  = card.select_one(".sqft, [class*='sqft'], .size, .sf")
                link     = card.select_one("a[href]")

                name  = name_el.get_text(strip=True)  if name_el  else ""
                addr  = addr_el.get_text(strip=True)  if addr_el  else ""
                price = _parse_price(price_el.get_text() if price_el else None)
                sqft  = _parse_sqft(size_el.get_text()  if size_el  else None)
                href  = link["href"] if link else ""
                if href and not href.startswith("http"):
                    href = urljoin("https://www.catylist.com", href)

                street, city, _, zp = _parse_address(addr)
                psf = round(price / sqft, 2) if (price and sqft and sqft > 0) else None

                if name or street:
                    listings.append(Listing(
                        listing_name   = name,
                        address        = street,
                        city           = city,
                        state          = "TX",
                        zip_code       = zp,
                        price          = price,
                        total_sqft     = sqft,
                        price_per_sqft = psf,
                        property_type  = _infer_prop_type(name + " " + addr),
                        listing_source = "Catylist",
                        listing_url    = href,
                    ))
            except Exception:
                continue

        time.sleep(REQUEST_DELAY)

    listings = [l for l in listings if _in_metro(l.city)]
    result.listings       = listings
    result.listings_found = len(listings)
    if not listings and not result.error:
        result.error = "No listings found"
    return result


# ─────────────────────────────────────────────────────────────────────────────
#  SCRAPER 5 — Land.com
# ─────────────────────────────────────────────────────────────────────────────

def scrape_land_com(session: requests.Session) -> ScraperResult:
    """Land.com — land marketplace. Paginate through San Antonio results."""
    result   = ScraperResult(source="Land.com")
    listings: list[Listing] = []

    base_urls = [
        "https://www.land.com/results/Texas/San-Antonio/",
        "https://www.land.com/results/Texas/Bexar-County/",
    ]

    for base_url in base_urls:
        for page in range(1, MAX_PAGES + 1):
            url = base_url if page == 1 else f"{base_url}{page}/"
            resp = _safe_get(session, url, result)
            if resp is None:
                break

            soup = BeautifulSoup(resp.text, "html.parser")

            listings.extend(_json_ld_listings(soup, "Land.com"))
            listings.extend(_next_data_listings(soup, "Land.com"))

            # Land.com uses .listing-card or .tile class
            cards = soup.select(
                ".listing-card, .listing-tile, .property-card, "
                ".listing, .result-card, article"
            )
            page_hits = 0
            for card in cards:
                try:
                    name_el  = card.select_one("h2, h3, .title, .listing-title, .name")
                    addr_el  = card.select_one(
                        ".listing-address, .address, .location, "
                        "[class*='address'], [class*='location']"
                    )
                    price_el = card.select_one(
                        ".listing-price, .price, [class*='price'], .asking"
                    )
                    size_el  = card.select_one(
                        ".listing-acres, .acreage, .acres, "
                        "[class*='acre'], [class*='size']"
                    )
                    link = card.select_one("a[href]")

                    name  = name_el.get_text(strip=True)  if name_el  else ""
                    addr  = addr_el.get_text(strip=True)  if addr_el  else ""
                    price = _parse_price(price_el.get_text() if price_el else None)
                    acres = _parse_acres(size_el.get_text()  if size_el  else None)
                    href  = link["href"] if link else ""
                    if href and not href.startswith("http"):
                        href = urljoin("https://www.land.com", href)

                    street, city, _, zp = _parse_address(addr)
                    psf = None
                    if price and acres and acres > 0:
                        psf = round(price / (acres * SQF_PER_ACRE), 2)

                    if name or street:
                        listings.append(Listing(
                            listing_name   = name,
                            address        = street,
                            city           = city,
                            state          = "TX",
                            zip_code       = zp,
                            price          = price,
                            lot_size_acres = acres,
                            price_per_sqft = psf,
                            property_type  = "land",
                            listing_source = "Land.com",
                            listing_url    = href,
                        ))
                        page_hits += 1
                except Exception:
                    continue

            time.sleep(REQUEST_DELAY)
            if page_hits == 0:
                break   # reached last page

        if listings:
            break   # first successful base_url is enough

    listings = [l for l in listings if _in_metro(l.city)]
    result.listings       = listings
    result.listings_found = len(listings)
    if not listings and not result.error:
        result.error = "No listings found"
    return result


# ─────────────────────────────────────────────────────────────────────────────
#  SCRAPER DISPATCH
# ─────────────────────────────────────────────────────────────────────────────

SCRAPER_FUNCS = {
    "Brevitas":           scrape_brevitas,
    "LandBrokerMLS":      scrape_landbrokermls,
    "CommercialExchange": scrape_commercialexchange,
    "Catylist":           scrape_catylist,
    "Land.com":           scrape_land_com,
}


# ─────────────────────────────────────────────────────────────────────────────
#  ADDRESS NORMALIZATION + DEDUPLICATION
# ─────────────────────────────────────────────────────────────────────────────

# Map verbose forms → short forms for address comparison
_ABBR: list[tuple[str, str]] = [
    (r"\bstreet\b",    "st"),
    (r"\broad\b",      "rd"),
    (r"\bavenue\b",    "ave"),
    (r"\bboulevard\b", "blvd"),
    (r"\blane\b",      "ln"),
    (r"\bdrive\b",     "dr"),
    (r"\bcourt\b",     "ct"),
    (r"\bplace\b",     "pl"),
    (r"\bparkway\b",   "pkwy"),
    (r"\bhighway\b",   "hwy"),
    (r"\bfreeway\b",   "fwy"),
    (r"\bnorth\b",     "n"),
    (r"\bsouth\b",     "s"),
    (r"\beast\b",      "e"),
    (r"\bwest\b",      "w"),
    (r"\bsuite\b",     "ste"),
    # Strip suite / unit numbers entirely
    (r"\b(ste|suite|unit|#)\s*[\w-]+", ""),
]


def _normalize_addr(addr: str) -> str:
    s = addr.lower().strip()
    for pat, rep in _ABBR:
        s = re.sub(pat, rep, s)
    s = re.sub(r"[^\w\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _addr_sim(a: str, b: str) -> float:
    return SequenceMatcher(None, _normalize_addr(a), _normalize_addr(b)).ratio()


def _prices_close(p1: Optional[float], p2: Optional[float]) -> bool:
    if not p1 or not p2:
        return True   # unknown price → can't exclude as dupe
    return abs(p1 - p2) / max(p1, p2) <= 0.05


def _completeness(lst: Listing) -> int:
    return sum(
        1 for f in vars(lst).values()
        if f not in (None, "", 0)
    )


def deduplicate(listings: list[Listing]) -> tuple[list[Listing], int]:
    """
    Normalise addresses, match on similarity ≥ 0.80 + price within 5%.
    Keep the most-complete record; append extra sources to also_listed_on.
    """
    kept:    list[Listing] = []
    removed: int           = 0

    for lst in listings:
        match = None
        for k in kept:
            if (lst.address and k.address and
                    _addr_sim(lst.address, k.address) >= 0.80 and
                    _prices_close(lst.price, k.price)):
                match = k
                break

        if match is None:
            kept.append(lst)
        else:
            removed += 1
            # Track cross-listed sources
            src = lst.listing_source
            if src and src != match.listing_source:
                existing = match.also_listed_on or ""
                if src not in existing:
                    match.also_listed_on = f"{existing}, {src}".lstrip(", ")

            # Prefer the more complete record as the keeper
            if _completeness(lst) > _completeness(match):
                idx = kept.index(match)
                # Merge sources into the new keeper
                lst.also_listed_on = match.also_listed_on
                lst.listing_source = match.listing_source
                kept[idx] = lst
            else:
                # Merge missing fields from duplicate into keeper
                for fname in vars(lst):
                    if fname in ("listing_source", "also_listed_on",
                                 "deal_type", "deal_type_reason", "flag"):
                        continue
                    if getattr(match, fname) in (None, "", 0):
                        setattr(match, fname, getattr(lst, fname))

    return kept, removed


# ─────────────────────────────────────────────────────────────────────────────
#  PSF CALCULATION
# ─────────────────────────────────────────────────────────────────────────────

def calculate_psf(lst: Listing) -> Listing:
    if lst.price is None or lst.price_per_sqft is not None:
        return lst
    sqft = lst.total_sqft
    if sqft is None and lst.lot_size_acres:
        sqft = lst.lot_size_acres * SQF_PER_ACRE
    if sqft and sqft > 0:
        lst.price_per_sqft = round(lst.price / sqft, 2)
    return lst


# ─────────────────────────────────────────────────────────────────────────────
#  CLAUDE BATCH CLASSIFICATION  (exactly ONE API call)
# ─────────────────────────────────────────────────────────────────────────────

_RULES = f"""
CLASSIFICATION RULES (current year = {CURRENT_YEAR}; apply first match):
1. DEVELOPMENT  — property_type is "land"
2. DEVELOPMENT  — year_built is null AND total_sqft is null
3. UNKNOWN      — year_built is null AND total_sqft is not null
4. UNKNOWN      — year_built is set AND ({CURRENT_YEAR} - year_built) < 10
5. VALUE-ADD    — year_built set, age ≥ 10, AND
                  (lot_size_acres is null OR total_sqft / (lot_size_acres * 43560) > 0.50)
6. DEVELOPMENT  — year_built set, age ≥ 10, lot_size_acres set,
                  total_sqft / (lot_size_acres * 43560) ≤ 0.50
7. UNKNOWN      — insufficient data

FLAGGING RULES:
  Compute separately:
    VA_avg   = mean price_per_sqft of VALUE-ADD listings that have price_per_sqft
    DEV_avg  = mean price_per_sqft of non-VALUE-ADD listings that have price_per_sqft
  Thresholds: VALUE-ADD = ${VA_PSF_THRESHOLD:.0f}/sqft,  all others = ${DEV_PSF_THRESHOLD:.0f}/sqft
  STRONG BUY : price_per_sqft ≤ threshold  AND  ≥ 10% below category average
  BELOW AVG  : price_per_sqft >  threshold  AND  ≥ 10% below category average
  WATCH      : price_per_sqft ≤ threshold  AND  < 10% below category average
  (blank)    : none of above, or price_per_sqft is null
"""


def classify_with_claude(client: anthropic.Anthropic,
                         listings: list[Listing]) -> list[Listing]:
    if not listings:
        return listings

    _use_api_call("classify")   # enforces MAX_API_CALLS guard

    batch = [
        {
            "index":          i,
            "listing_name":   l.listing_name,
            "address":        l.address,
            "city":           l.city,
            "price":          l.price,
            "total_sqft":     l.total_sqft,
            "price_per_sqft": l.price_per_sqft,
            "year_built":     l.year_built,
            "lot_size_acres": l.lot_size_acres,
            "property_type":  l.property_type,
            "source":         l.listing_source,
        }
        for i, l in enumerate(listings)
    ]

    prompt = f"""You are a commercial real estate analyst.

{_RULES}

LISTINGS:
{json.dumps(batch, indent=2)}

Return ONLY a valid JSON array — no markdown, no explanation:
[
  {{"index": 0, "deal_type": "VALUE-ADD", "deal_type_reason": "40-yr structure covers 62% of site", "flag": "WATCH"}},
  {{"index": 1, "deal_type": "DEVELOPMENT", "deal_type_reason": "vacant land parcel", "flag": "STRONG BUY"}},
  ...
]"""

    t0 = time.time()
    response = client.messages.create(
        model=MODEL,
        max_tokens=8000,
        messages=[{"role": "user", "content": prompt}],
    )
    print(f"  Claude responded in {time.time() - t0:.1f}s")

    text = response.content[0].text if response.content else ""

    arr: Optional[list] = None
    m = re.search(r"```(?:json)?\s*(\[[\s\S]*?\])\s*```", text)
    if m:
        try:
            arr = json.loads(m.group(1))
        except Exception:
            pass
    if arr is None:
        start = text.find("[")
        if start >= 0:
            depth = 0
            for i, ch in enumerate(text[start:], start):
                if ch == "[":
                    depth += 1
                elif ch == "]":
                    depth -= 1
                    if depth == 0:
                        try:
                            arr = json.loads(text[start : i + 1])
                        except Exception:
                            pass
                        break

    if not arr:
        print("  ⚠ Could not parse Claude's classification response")
        return listings

    for item in arr:
        try:
            idx = int(item["index"])
        except (KeyError, ValueError, TypeError):
            continue
        if 0 <= idx < len(listings):
            listings[idx].deal_type        = item.get("deal_type", "UNKNOWN")
            listings[idx].deal_type_reason = item.get("deal_type_reason", "")
            listings[idx].flag             = item.get("flag", "")

    return listings


# ─────────────────────────────────────────────────────────────────────────────
#  EXCEL EXPORT
# ─────────────────────────────────────────────────────────────────────────────

COLUMNS = [
    ("Listing Name",       "listing_name",      28),
    ("Address",            "address",           24),
    ("City",               "city",              14),
    ("State",              "state",              6),
    ("Zip Code",           "zip_code",           9),
    ("Price",              "price",             14),
    ("Total Sqft",         "total_sqft",        12),
    ("Price / Sqft",       "price_per_sqft",    12),
    ("Year Built",         "year_built",        10),
    ("Lot Size (Acres)",   "lot_size_acres",    14),
    ("Zoning",             "zoning",            10),
    ("Days on Market",     "days_on_market",    13),
    ("Property Type",      "property_type",     18),
    ("Source",             "listing_source",    14),
    ("Also Listed On",     "also_listed_on",    16),
    ("Listing URL",        "listing_url",       42),
    ("Deal Type",          "deal_type",         13),
    ("Deal Type Reason",   "deal_type_reason",  46),
    ("FLAG",               "flag",              12),
]

C_HDR  = "1F3864"
C_VA   = "E2EFDA"
C_DEV  = "DDEBF7"
C_UNK  = "FFF2CC"
C_SB   = "C00000"
C_BA   = "ED7D31"
C_WA   = "FFE699"
T_SIDE = Side(style="thin", color="D9D9D9")
T_BORD = Border(left=T_SIDE, right=T_SIDE, top=T_SIDE, bottom=T_SIDE)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def export_excel(listings: list[Listing], filepath: str,
                 results: list[ScraperResult], dupes: int) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Listings"

    hf = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    for ci, (hdr, _, w) in enumerate(COLUMNS, 1):
        c            = ws.cell(row=1, column=ci, value=hdr)
        c.font       = hf
        c.fill       = _fill(C_HDR)
        c.alignment  = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border     = T_BORD
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes  = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    url_ci = next(i for i, (_, f, _) in enumerate(COLUMNS, 1) if f == "listing_url")

    for ri, lst in enumerate(listings, 2):
        row_bg = {"VALUE-ADD": C_VA, "DEVELOPMENT": C_DEV}.get(lst.deal_type, C_UNK)
        for ci, (_, fname, _) in enumerate(COLUMNS, 1):
            val = getattr(lst, fname, None)
            if fname == "price"                                          and val is not None: val = float(val)
            elif fname in ("total_sqft", "year_built", "days_on_market") and val is not None: val = int(val)
            elif fname in ("price_per_sqft", "lot_size_acres")           and val is not None: val = float(val)

            c           = ws.cell(row=ri, column=ci, value=val)
            c.fill      = _fill(row_bg)
            c.border    = T_BORD
            c.alignment = Alignment(vertical="top", wrap_text=True)

            if   fname == "price"          and val is not None: c.number_format = '$#,##0'
            elif fname == "price_per_sqft" and val is not None: c.number_format = '$#,##0.00'
            elif fname == "lot_size_acres" and val is not None: c.number_format = '0.00'
            elif fname == "total_sqft"     and val is not None: c.number_format = '#,##0'

            if fname == "flag" and val in ("STRONG BUY", "BELOW AVG", "WATCH"):
                colors = {"STRONG BUY": (C_SB, "FFFFFF"), "BELOW AVG": (C_BA, "000000"), "WATCH": (C_WA, "000000")}
                bg, fg = colors[val]
                c.fill      = _fill(bg)
                c.font      = Font(name="Calibri", bold=True, color=fg, size=10)
                c.alignment = Alignment(horizontal="center", vertical="center")

            if ci == url_ci and val and str(val).startswith("http"):
                c.hyperlink = str(val)
                c.font      = Font(name="Calibri", color="0563C1", underline="single", size=10)

        ws.row_dimensions[ri].height = 48

    # Summary sheet
    ss = wb.create_sheet("Summary")
    ss.column_dimensions["A"].width = 32
    ss.column_dimensions["B"].width = 20

    def sh(row: int, text: str) -> None:
        c           = ss.cell(row=row, column=1, value=text)
        c.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        c.fill      = _fill(C_HDR)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ss.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ss.row_dimensions[row].height = 20

    def sr(row: int, label: str, value) -> None:
        a = ss.cell(row=row, column=1, value=label)
        b = ss.cell(row=row, column=2, value=value)
        a.font      = Font(name="Calibri", size=10)
        b.font      = Font(name="Calibri", size=10, bold=True)
        b.alignment = Alignment(horizontal="right")
        for x in (a, b):
            x.border = T_BORD

    total = len(listings)
    va  = sum(1 for l in listings if l.deal_type == "VALUE-ADD")
    dev = sum(1 for l in listings if l.deal_type == "DEVELOPMENT")
    unk = sum(1 for l in listings if l.deal_type == "UNKNOWN")
    sb  = sum(1 for l in listings if l.flag == "STRONG BUY")
    ba  = sum(1 for l in listings if l.flag == "BELOW AVG")
    wa  = sum(1 for l in listings if l.flag == "WATCH")

    r = 1
    sh(r, "PropScout V2 Phase 1 — SA Metro Summary");       r += 1
    sh(r, "Run info");                                       r += 1
    sr(r, "Report date", datetime.now().strftime("%Y-%m-%d %H:%M")); r += 1
    sr(r, "Geography",   "San Antonio TX metro + suburbs"); r += 1
    sh(r, "Listing counts");                                 r += 1
    sr(r, "Total unique listings", total);                   r += 1
    sr(r, "Duplicates removed",    dupes);                   r += 1
    sh(r, "Deal type breakdown");                            r += 1
    sr(r, "VALUE-ADD",   va);  r += 1
    sr(r, "DEVELOPMENT", dev); r += 1
    sr(r, "UNKNOWN",     unk); r += 1
    sh(r, "Flagged listings");                               r += 1
    sr(r, "STRONG BUY",  sb);  r += 1
    sr(r, "BELOW AVG",   ba);  r += 1
    sr(r, "WATCH",       wa);  r += 1
    sh(r, "Source diagnostics");                             r += 1
    for res in results:
        notes = res.error or "OK"
        sr(r, res.source, f"{res.listings_found} listings  [{notes}]"); r += 1

    wb.save(filepath)


# ─────────────────────────────────────────────────────────────────────────────
#  SOURCE DIAGNOSTIC TABLE
# ─────────────────────────────────────────────────────────────────────────────

def print_diagnostics(results: list[ScraperResult], dupes: int,
                      total_after: int) -> None:
    W = 72
    print()
    print("─" * W)
    print("  SOURCE DIAGNOSTIC TABLE")
    print("─" * W)
    print(f"  {'Source':<22} {'HTTP':>6}  {'Listings':>9}  {'Notes'}")
    print(f"  {'─'*22} {'─'*6}  {'─'*9}  {'─'*28}")
    for r in results:
        status = str(r.http_status) if r.http_status else "—"
        notes  = r.error or "OK"
        src_label = r.source
        if not SOURCES.get(r.source, True):
            src_label = f"{r.source} [disabled]"
            notes = "disabled"
        print(f"  {src_label:<22} {status:>6}  {r.listings_found:>9}  {notes[:38]}")
    print("─" * W)
    total_raw = sum(r.listings_found for r in results)
    print(f"  Raw total: {total_raw}   Dupes removed: {dupes}   Unique: {total_after}")
    print("─" * W)
    print()


# ─────────────────────────────────────────────────────────────────────────────
#  TERMINAL SUMMARY
# ─────────────────────────────────────────────────────────────────────────────

def print_summary(listings: list[Listing], dupes: int,
                  excel_path: str, json_path: str) -> None:
    W  = 62
    hr = lambda c="─": print(c * W)

    total = len(listings)
    va  = sum(1 for l in listings if l.deal_type == "VALUE-ADD")
    dev = sum(1 for l in listings if l.deal_type == "DEVELOPMENT")
    unk = sum(1 for l in listings if l.deal_type == "UNKNOWN")
    sb  = sum(1 for l in listings if l.flag == "STRONG BUY")
    ba  = sum(1 for l in listings if l.flag == "BELOW AVG")
    wa  = sum(1 for l in listings if l.flag == "WATCH")

    print(); hr("═")
    print("  PROPSCOUT V2 — FINAL SUMMARY".center(W))
    hr("═")
    print(f"  Total unique listings:   {total}")
    print(f"  Duplicates removed:      {dupes}")
    hr()
    print("  DEAL TYPE BREAKDOWN")
    print(f"    VALUE-ADD:             {va}")
    print(f"    DEVELOPMENT:           {dev}")
    print(f"    UNKNOWN:               {unk}")
    hr()
    print("  FLAGGED LISTINGS")
    print(f"    STRONG BUY:            {sb}")
    print(f"    BELOW AVG:             {ba}")
    print(f"    WATCH:                 {wa}")
    hr()
    if sb or ba or wa:
        print("  FLAGGED LISTING DETAILS")
        for l in listings:
            if l.flag:
                addr = f"{l.address}, {l.city}"[:40]
                psf  = f"${l.price_per_sqft:.2f}/sqft" if l.price_per_sqft else "no PSF"
                print(f"    [{l.flag:<10}] {addr:<42} {psf}")
        hr()
    hr("═")
    print(f"  Excel  → {excel_path}")
    print(f"  JSON   → {json_path}")
    hr("═"); print()


# ─────────────────────────────────────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser(
        description="PropScout V2 Phase 1 — SA Metro scraper + classifier"
    )
    parser.add_argument(
        "--dry-run", action="store_true",
        help="Scrape and deduplicate only; skip the Claude API call (free test)",
    )
    args = parser.parse_args()

    W  = 62
    hr = lambda c="─": print(c * W)

    print(); hr("═")
    print("  PropScout V2 Phase 1 — SA Metro Scraper".center(W))
    hr("═")
    enabled_sources = [s for s, on in SOURCES.items() if on]
    print(f"  Sources:   {', '.join(enabled_sources)}")
    print(f"  API calls: {'0 (--dry-run)' if args.dry_run else '1 (Claude classify)'}")
    print(f"  Output:    {EXCEL_PATH}  ·  {RAW_JSON}")
    hr(); print()

    if not os.environ.get("ANTHROPIC_API_KEY") and not args.dry_run:
        sys.exit("ERROR: ANTHROPIC_API_KEY not set. Set it or use --dry-run.")

    session = requests.Session()
    session.headers.update({
        "User-Agent":                USER_AGENT,
        "Accept":                    "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language":           "en-US,en;q=0.9",
        "Accept-Encoding":           "gzip, deflate, br",
        "Connection":                "keep-alive",
        "Upgrade-Insecure-Requests": "1",
        "Sec-Fetch-Dest":            "document",
        "Sec-Fetch-Mode":            "navigate",
        "Sec-Fetch-Site":            "none",
        "Cache-Control":             "max-age=0",
    })

    # ── Step 1: Scrape all enabled sources ────────────────────────────────────
    all_raw:     list[Listing]       = []
    diag_results: list[ScraperResult] = []

    for source_name, enabled in SOURCES.items():
        if not enabled:
            diag_results.append(ScraperResult(source=source_name, error="disabled"))
            continue

        print(f"  ── {source_name} ──────────────────────────────────")
        scraper = SCRAPER_FUNCS[source_name]
        try:
            result = scraper(session)
        except Exception as exc:
            result = ScraperResult(source=source_name, error=f"Unhandled: {exc}")
            print(f"  ✗ {exc}")

        diag_results.append(result)
        print(f"  → {result.listings_found} listing(s)"
              + (f"  [{result.error}]" if result.error else ""))
        all_raw.extend(result.listings)
        print()

    # ── Step 2: PSF calculation ───────────────────────────────────────────────
    all_raw = [calculate_psf(l) for l in all_raw]

    # ── Step 3: Deduplicate ───────────────────────────────────────────────────
    listings, dupes = deduplicate(all_raw)
    print(f"  Dedup: {dupes} removed  |  {len(listings)} unique listings")
    hr()

    # ── Step 4: Save raw JSON (before Claude) ─────────────────────────────────
    raw_dicts = []
    for l in listings:
        d = {k: v for k, v in vars(l).items() if v not in (None, "", 0)}
        raw_dicts.append(d)
    with open(RAW_JSON, "w", encoding="utf-8") as fh:
        json.dump(raw_dicts, fh, indent=2, default=str)
    print(f"  Raw JSON saved → {RAW_JSON}  ({len(raw_dicts)} records)")
    hr()

    if not listings:
        print("  No listings scraped — all sources blocked or returned nothing.")
        print_diagnostics(diag_results, dupes, 0)
        sys.exit(0)

    # ── Step 5: Claude classification (skipped in --dry-run) ─────────────────
    if args.dry_run:
        print("  --dry-run: skipping Claude API call")
        hr()
    else:
        client   = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        print(f"  Classifying {len(listings)} listings with Claude …")
        listings = classify_with_claude(client, listings)
        va  = sum(1 for l in listings if l.deal_type == "VALUE-ADD")
        dev = sum(1 for l in listings if l.deal_type == "DEVELOPMENT")
        unk = sum(1 for l in listings if l.deal_type == "UNKNOWN")
        print(f"  VALUE-ADD={va}  DEVELOPMENT={dev}  UNKNOWN={unk}")
        hr()

    # ── Step 6: Export Excel ──────────────────────────────────────────────────
    print(f"  Exporting → {EXCEL_PATH} …")
    export_excel(listings, EXCEL_PATH, diag_results, dupes)

    # ── Diagnostic table + summary ────────────────────────────────────────────
    print_diagnostics(diag_results, dupes, len(listings))
    if not args.dry_run:
        print_summary(listings, dupes, EXCEL_PATH, RAW_JSON)


if __name__ == "__main__":
    main()
